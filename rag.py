#!/usr/bin/env python3
import os
import hashlib
import re
import json
import shutil
import importlib
from datetime import datetime
import requests
from pathlib import Path
import numpy as np
from tqdm import tqdm
from dotenv import load_dotenv

load_dotenv()

from backend.config import settings

HF_TOKEN = settings.HF_TOKEN or ""
GROQ_API_KEY = os.getenv("GROQ_API_KEY", "")

# -------- Retrieval --------
import faiss
from rank_bm25 import BM25Okapi

from sqlalchemy.orm import Session

from backend.database import SessionLocal
from backend.models import Message as MessageModel

# =====================================
# CONFIG
# =====================================
DATA_DIR = "data"
UPLOADS_DIR = "uploads"
INDEX_DIR = "index_data"
CHUNK_DIR = f"{INDEX_DIR}/chunks"

# HF Embedding API
HF_EMB_MODEL = "sentence-transformers/all-MiniLM-L6-v2"
HF_EMB_URL = "https://api-inference.huggingface.co/models/sentence-transformers/all-MiniLM-L6-v2"

# Groq LLM API
GROQ_MODEL = "llama-3.3-70b-versatile"
GROQ_FALLBACK_MODEL = "llama-3.1-8b-instant"
GROQ_API_URL = "https://api.groq.com/openai/v1/chat/completions"

CHUNK_SIZE = 3000
OVERLAP = 300
TOP_BM25 = 20
TOP_DENSE = 20
FINAL_TOPK = 5

_startup_reindex_checked = False

os.makedirs(CHUNK_DIR, exist_ok=True)
os.makedirs(DATA_DIR, exist_ok=True)

print("✅ LegalLens RAG Ready (API-based, no local models)")


def _clear_directory_contents(directory_path: str) -> None:
    if not os.path.exists(directory_path):
        return

    for entry_name in os.listdir(directory_path):
        entry_path = os.path.join(directory_path, entry_name)
        if os.path.isdir(entry_path) and not os.path.islink(entry_path):
            shutil.rmtree(entry_path)
        else:
            os.remove(entry_path)


def _clear_index_artifacts() -> None:
    os.makedirs(INDEX_DIR, exist_ok=True)
    os.makedirs(CHUNK_DIR, exist_ok=True)
    _clear_directory_contents(CHUNK_DIR)

    for artifact_name in ["faiss.index", "bm25.json", "meta.json"]:
        artifact_path = os.path.join(INDEX_DIR, artifact_name)
        if os.path.exists(artifact_path):
            os.remove(artifact_path)


def clear_index() -> None:
    _clear_directory_contents(INDEX_DIR)
    _clear_directory_contents(DATA_DIR)
    os.makedirs(CHUNK_DIR, exist_ok=True)
    os.makedirs(DATA_DIR, exist_ok=True)


def clean_text(text: str) -> str:
    return text.encode("utf-8", errors="replace").decode("utf-8")


def read_text_utf8(path: str | Path) -> str:
    with open(path, "r", encoding="utf-8") as file:
        return file.read()


def write_text_utf8(path: str | Path, text: str) -> None:
    with open(path, "w", encoding="utf-8") as file:
        file.write(clean_text(text))


def load_json_utf8(path: str | Path):
    with open(path, "r", encoding="utf-8") as file:
        return json.load(file)


def dump_json_utf8(path: str | Path, data) -> None:
    with open(path, "w", encoding="utf-8") as file:
        json.dump(data, file, ensure_ascii=False)


def extract_pdf_text(path: Path) -> str:
    try:
        PdfReader = importlib.import_module("pypdf").PdfReader

        with open(path, "rb") as file:
            reader = PdfReader(file)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        return clean_text(text)
    except Exception as exc:
        print(f"⚠️  Could not read PDF {path.name}: {exc}")
        return ""


def extract_text_from_source(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return clean_text(extract_docx(path))
    if suffix == ".pdf":
        return extract_pdf_text(path)
    if suffix in {".txt", ".md"}:
        try:
            with open(path, "r", encoding="utf-8") as file:
                return clean_text(file.read())
        except Exception as exc:
            print(f"⚠️  Could not read text file {path.name}: {exc}")
            return ""
    return ""

# =====================================
# HF EMBEDDING API
# =====================================
def get_hf_embeddings_batch(texts: list) -> list | None:
    headers = {
        "Authorization": f"Bearer {HF_TOKEN}",
        "Content-Type": "application/json"
    }
    all_embeddings = []
    batch_size = 8
    try:
        for i in range(0, len(texts), batch_size):
            batch = texts[i:i + batch_size]
            response = requests.post(
                HF_EMB_URL,
                headers=headers,
                json={"inputs": batch},
                timeout=30
            )
            if response.status_code == 200:
                all_embeddings.extend(response.json())
            else:
                print(f"⚠️ HF embedding failed: {response.status_code}")
                return None
        return all_embeddings
    except Exception as e:
        print(f"⚠️ HF embedding exception: {e}")
        return None


def hash_embed(text: str, dim: int = 384) -> list:
    vec = []
    for i in range(dim):
        h = hashlib.md5(f"{text}__{i}".encode()).hexdigest()
        vec.append(int(h[:8], 16) / 0xffffffff - 0.5)
    arr = np.array(vec)
    norm = np.linalg.norm(arr)
    if norm > 0:
        arr = arr / norm
    return arr.tolist()


def fallback_embeddings(texts: list) -> np.ndarray:
    """Simple hash-based fallback embeddings when API unavailable."""
    embeddings = [hash_embed(text) for text in texts]
    return np.array(embeddings, dtype=np.float32)


# =====================================
# GROQ LLM API
# =====================================
def call_llm(messages: list, system: str = "") -> str:
    """Call Groq API for fast LLM inference."""
    if isinstance(messages, dict):
        messages = [messages]

    sanitized_messages = []
    for msg in messages:
        if not isinstance(msg, dict):
            continue
        role = msg.get("role")
        content = msg.get("content")
        if role in {"user", "assistant"} and content:
            sanitized_messages.append({"role": role, "content": str(content)})

    headers = {
        "Authorization": f"Bearer {GROQ_API_KEY}",
        "Content-Type": "application/json"
    }

    if system:
        sanitized_messages = [{"role": "system", "content": system}] + sanitized_messages

    payload = {
        "model": GROQ_MODEL,
        "messages": sanitized_messages,
        "max_tokens": 3000,
        "temperature": 0.2
    }

    try:
        response = requests.post(
            GROQ_API_URL,
            headers=headers,
            json=payload,
            timeout=30
        )

        if response.status_code == 200:
            return response.json()["choices"][0]["message"]["content"].strip()

        if GROQ_MODEL != GROQ_FALLBACK_MODEL:
            print(
                f"⚠️ Groq primary model failed: status={response.status_code}, body={response.text}"
            )
            fallback_payload = dict(payload)
            fallback_payload["model"] = GROQ_FALLBACK_MODEL
            fallback_response = requests.post(
                GROQ_API_URL,
                headers=headers,
                json=fallback_payload,
                timeout=30
            )
            if fallback_response.status_code == 200:
                return fallback_response.json()["choices"][0]["message"]["content"].strip()
            print(
                f"⚠️ Groq fallback model failed: status={fallback_response.status_code}, body={fallback_response.text}"
            )
            return call_hf_fallback(sanitized_messages)

        print(f"⚠️ Groq failed: status={response.status_code}, body={response.text}")
        return call_hf_fallback(sanitized_messages)

    except requests.exceptions.Timeout:
        return "Request timed out. Please try again."
    except Exception as e:
        return f"Error: {str(e)}"


def call_hf_fallback(messages: list) -> str:
    """Fallback to HF featherless-ai if Groq unavailable."""
    headers = {"Authorization": f"Bearer {HF_TOKEN}"}
    sanitized_messages = []
    for msg in messages:
        if not isinstance(msg, dict):
            continue
        role = msg.get("role")
        content = msg.get("content")
        if role in {"system", "user", "assistant"} and content:
            sanitized_messages.append({"role": role, "content": str(content)})

    payload = {
        "model": "mistralai/Mistral-7B-Instruct-v0.2",
        "messages": sanitized_messages,
        "max_tokens": 3000,
        "temperature": 0.2
    }
    try:
        response = requests.post(
            "https://router.huggingface.co/featherless-ai/v1/chat/completions",
            headers=headers,
            json=payload,
            timeout=120
        )
        if response.status_code == 200:
            return response.json()["choices"][0]["message"]["content"].strip()
        return f"API Error {response.status_code}: {response.text[:200]}"
    except Exception as e:
        return f"Error: {str(e)}"


# =====================================
# PROMPT BUILDER
# =====================================
def sanitize_history(messages: list) -> list:
    cleaned = []

    def message_timestamp(value):
        if isinstance(value, datetime):
            return value
        if isinstance(value, str):
            try:
                return datetime.fromisoformat(value.replace("Z", "+00:00"))
            except Exception:
                return datetime.min
        return datetime.min

    prepared = []
    for msg in messages:
        if not isinstance(msg, dict):
            continue
        role = msg.get("role")
        content = msg.get("content")
        if role not in {"user", "assistant"} or not content:
            continue
        prepared.append({
            "role": role,
            "content": str(content),
            "timestamp": msg.get("timestamp"),
        })

    if any(item.get("timestamp") is not None for item in prepared):
        prepared.sort(key=lambda item: message_timestamp(item.get("timestamp")))

    for msg in prepared:
        if cleaned and cleaned[-1]["role"] == msg["role"]:
            cleaned[-1] = {"role": msg["role"], "content": msg["content"]}
        else:
            cleaned.append({"role": msg["role"], "content": msg["content"]})

    while cleaned and cleaned[0]["role"] != "user":
        cleaned.pop(0)

    return cleaned


def build_prompt(question: str, context: str, chat_history: list) -> tuple[list, str]:
    system = (
        "You are a legal AI assistant. Answer questions based on the provided context. "
        "Be concise and direct. Do not output JSON. Do not repeat the question. "
        "If the context does not contain the answer, say so clearly."
    )
    history = sanitize_history(chat_history)[-6:]
    current_user_message = {
        "role": "user",
        "content": f"Context:\n{context}\n\nQuestion:\n{question}",
    }

    return history + [current_user_message], system


# =====================================
# DOCX EXTRACTION
# =====================================
def extract_docx(path: Path) -> str:
    try:
        import docx
        doc = docx.Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    paragraphs.append(row_text)
        return "\n".join(paragraphs)
    except Exception as e:
        print(f"⚠️  Could not read {path.name}: {e}")
        return ""


def check_and_reindex_corrupt_chunks() -> None:
    global _startup_reindex_checked
    if _startup_reindex_checked:
        return
    _startup_reindex_checked = True

    chunk_dir = Path(CHUNK_DIR)
    if not chunk_dir.exists():
        return

    for chunk_file in chunk_dir.rglob("*.txt"):
        try:
            with open(chunk_file, "r", encoding="utf-8") as file:
                if "Â" in file.read():
                    print("⚠️ Corrupt chunks detected, re-indexing...")
                    clear_index()
                    index_documents()
                    return
        except Exception:
            continue


# =====================================
# SYNC UPLOADS → DATA
# =====================================
def sync_uploads_to_data():
    uploads_dir = Path(UPLOADS_DIR)
    data_dir = Path(DATA_DIR)
    index_dir = Path(INDEX_DIR)
    chunk_dir = Path(CHUNK_DIR)
    data_dir.mkdir(exist_ok=True)
    index_dir.mkdir(exist_ok=True)
    chunk_dir.mkdir(parents=True, exist_ok=True)

    if not uploads_dir.exists():
        return False

    supported = {".txt", ".md", ".pdf", ".docx"}
    desired_data_files = set()
    changed = False

    for f in uploads_dir.rglob("*"):
        if not f.is_file():
            continue
        if f.suffix.lower() not in supported:
            continue

        if f.suffix.lower() == ".docx":
            dest = data_dir / (f.stem + "_docx.txt")
            desired_data_files.add(dest.name)
            if not dest.exists():
                text = extract_text_from_source(f)
                if text.strip():
                    write_text_utf8(dest, text)
                    print(f"📥 Extracted & synced: {f.name} → {dest.name}")
                    changed = True
        elif f.suffix.lower() == ".pdf":
            dest = data_dir / (f.stem + "_pdf.txt")
            desired_data_files.add(dest.name)
            if not dest.exists():
                text = extract_text_from_source(f)
                if text.strip():
                    write_text_utf8(dest, text)
                    print(f"📥 Extracted & synced: {f.name} → {dest.name}")
                    changed = True
        else:
            dest = data_dir / f.name
            desired_data_files.add(dest.name)
            if not dest.exists():
                text = extract_text_from_source(f)
                if text.strip():
                    write_text_utf8(dest, text)
                    print(f"📥 Synced: {f.name}")
                    changed = True

    for existing_file in data_dir.rglob("*"):
        if not existing_file.is_file():
            continue
        if existing_file.name not in desired_data_files:
            existing_file.unlink()
            changed = True

    meta_path = Path(INDEX_DIR) / "meta.json"
    if meta_path.exists():
        try:
            meta = load_json_utf8(meta_path)
        except Exception:
            meta = []

        indexed_files = {
            item.get("doc")
            for item in meta
            if isinstance(item, dict) and item.get("doc")
        }
        stale_index_files = indexed_files - desired_data_files
        if stale_index_files:
            for item in meta:
                if not isinstance(item, dict):
                    continue
                if item.get("doc") in stale_index_files:
                    chunk_path = item.get("path")
                    if chunk_path and os.path.exists(chunk_path):
                        os.remove(chunk_path)

            _clear_index_artifacts()
            changed = True

    return changed


# =====================================
# INDEXING
# =====================================
def chunk_text(text):
    i = 0
    while i < len(text):
        yield text[i:i + CHUNK_SIZE].strip()
        i += CHUNK_SIZE - OVERLAP


def index_documents():
    print("📚 Indexing documents...")
    sync_uploads_to_data()
    _clear_index_artifacts()

    files = list(Path(DATA_DIR).rglob("*"))
    meta, corpus, vectors = [], [], []
    cid = 0

    for f in tqdm(files):
        if not f.is_file():
            continue

        suffix = f.suffix.lower()
        text = extract_text_from_source(f)

        if not text.strip():
            continue

        for chunk in chunk_text(text):
            if not chunk:
                continue

            chunk_path = Path(CHUNK_DIR) / f"chunk_{cid}.txt"
            write_text_utf8(chunk_path, chunk)

            meta.append({
                "id": cid,
                "doc": f.name,
                "path": str(chunk_path)
            })

            corpus.append(chunk)
            vectors.append(chunk)
            cid += 1

    if cid == 0:
        print("⚠️  No documents found to index.")
        return 0

    print("🔢 Computing embeddings via HF API...")
    embeddings = get_hf_embeddings_batch(vectors)
    if embeddings is None:
        embeddings = fallback_embeddings(vectors)
    else:
        embeddings = np.array(embeddings, dtype=np.float32)
    index = faiss.IndexFlatL2(embeddings.shape[1])
    index.add(embeddings.astype(np.float32))

    faiss.write_index(index, f"{INDEX_DIR}/faiss.index")
    dump_json_utf8(f"{INDEX_DIR}/bm25.json", corpus)
    dump_json_utf8(f"{INDEX_DIR}/meta.json", meta)

    print(f"✅ Indexed {cid} chunks")
    return cid


# =====================================
# ANSWER PARSER
# =====================================
def parse_answer(raw: str) -> str:
    raw = raw.strip()

    for prefix in ["Answer:", "answer:", "\nAnswer:"]:
        if raw.startswith(prefix):
            raw = raw[len(prefix):].strip()
            break

    if raw.startswith("Output:"):
        raw = raw[len("Output:"):].strip()

    try:
        data = json.loads(raw)
        if isinstance(data, dict) and "answer" in data:
            return data["answer"].strip()
    except Exception:
        pass

    match = re.search(r'"answer"\s*:\s*"(.*?)"(?:\s*,|\s*\})', raw, re.DOTALL)
    if match:
        return match.group(1).strip()

    return raw


# =====================================
# HYBRID RETRIEVER
# =====================================
def hybrid_retrieve(query):
    data_changed = sync_uploads_to_data()

    if not Path(f"{INDEX_DIR}/faiss.index").exists():
        index_documents()

    if not Path(f"{INDEX_DIR}/faiss.index").exists():
        return "No documents have been indexed yet. Please upload a document first."

    # Re-index if the uploads/data set changed or the index no longer matches it
    data_files = set(f.name for f in Path(DATA_DIR).rglob("*") if f.is_file())
    meta_path = Path(f"{INDEX_DIR}/meta.json")
    if meta_path.exists():
        meta = load_json_utf8(meta_path)
        indexed_files = set(m["doc"] for m in meta)
        if data_changed or data_files != indexed_files:
            print("🔄 New files detected, re-indexing...")
            index_documents()

    index = faiss.read_index(f"{INDEX_DIR}/faiss.index")
    corpus = load_json_utf8(f"{INDEX_DIR}/bm25.json")
    meta = load_json_utf8(f"{INDEX_DIR}/meta.json")

    if not corpus:
        return "No documents have been indexed yet."

    # BM25 retrieval
    bm25 = BM25Okapi([c.split() for c in corpus])
    bm25_ids = np.argsort(
        bm25.get_scores(query.split())
    )[::-1][:TOP_BM25]

    # Dense retrieval via HF embedding API
    q_emb = get_hf_embeddings_batch([query])
    if q_emb is None:
        q_emb = fallback_embeddings([query])
    else:
        q_emb = np.array(q_emb, dtype=np.float32)
    _, dense_ids = index.search(q_emb.astype(np.float32), TOP_DENSE)

    candidates = list(set(bm25_ids.tolist() + dense_ids[0].tolist()))

    # Simple reranking by BM25 score (no CrossEncoder needed)
    bm25_scores = bm25.get_scores(query.split())
    ranked = sorted(candidates, key=lambda x: bm25_scores[x], reverse=True)[:FINAL_TOPK]

    docs = []
    for cid in ranked:
        m = meta[cid]
        content = read_text_utf8(m["path"])
        docs.append(clean_text(content))

    return "\n\n".join(docs[:3])


# =====================================
# MEMORY STORE
# =====================================
def get_session_history(session_id: str) -> list[dict]:
    db: Session = SessionLocal()
    try:
        rows = (
            db.query(MessageModel)
            .filter(MessageModel.session_id == session_id)
            .order_by(MessageModel.timestamp.asc())
            .all()
        )
        return [
            {"role": row.role, "content": row.content, "timestamp": row.timestamp}
            for row in rows
            if row.role in {"user", "assistant"}
        ]
    finally:
        db.close()


check_and_reindex_corrupt_chunks()


# =====================================
# MAIN CHAIN
# =====================================
class RAGChain:
    def invoke(self, inputs: dict, config: dict = None) -> str:
        question = inputs["question"]
        session_id = "default"

        if config and "configurable" in config:
            session_id = config["configurable"].get("session_id", "default")

        history = get_session_history(session_id)
        context = hybrid_retrieve(question)
        prompt, system = build_prompt(question, context, history)
        raw_answer = call_llm(prompt, system=system)
        answer = parse_answer(raw_answer)

        return answer


chain_with_memory = RAGChain()


# =====================================
# MAIN LOOP
# =====================================
def main():
    print("🧠 LegalLens RAG Ready (Groq + HF Embeddings)\n")
    session_id = "legal-session"

    while True:
        query = input(">> ")
        if query.lower() == "exit":
            break

        answer = chain_with_memory.invoke(
            {"question": query},
            config={"configurable": {"session_id": session_id}}
        )

        print("\n📄 RESPONSE\n")
        print(answer)
        print()


if __name__ == "__main__":
    main()